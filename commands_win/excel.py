"""Excel commands for the Windows build — read, write, append rows, nicknames.

Nicknames let you say "my budget sheet" instead of a full path:
register once, then use the nickname in every command.
"""
import csv
import io
import json
import os
import re
import zipfile

from common import data_dir

try:
    import openpyxl
    _HAVE_OPENPYXL = True
except ImportError:
    _HAVE_OPENPYXL = False

ALIAS_FILE = os.path.join(data_dir(), "excel_aliases.json")
_MAX_ROWS = 50


def _aliases():
    try:
        with open(ALIAS_FILE) as f:
            d = json.load(f)
            return d if isinstance(d, dict) else {}
    except (OSError, ValueError):
        return {}


def _save_aliases(a):
    with open(ALIAS_FILE, "w") as f:
        json.dump(a, f, indent=1)


def _resolve(ref):
    """Nickname -> path, else a literal path. Returns (path, error)."""
    ref = (ref or "").strip().strip("\"'")
    if not ref:
        return "", "Tell me which workbook — a registered nickname or a .xlsx path."
    for nick, path in _aliases().items():
        if nick.lower() == ref.lower():
            return path, ""
    p = os.path.expanduser(ref)
    if not p.lower().endswith(".xlsx"):
        return "", (f"No sheet registered as {ref!r}, and it isn't a .xlsx path. "
                    "Register one with: remember sheet <nickname> <path>")
    if not os.path.isfile(p):
        return "", f"Workbook not found: {p}"
    return p, ""


def _sanitize_xlsx(path):
    """Strip empty <v/> tags openpyxl writes for formula cells.

    Without this, Excel shows a "problem with some content" repair prompt.
    """
    with zipfile.ZipFile(path) as zin:
        items = {n: zin.read(n) for n in zin.namelist()}
    changed = False
    for name, data in items.items():
        if name.endswith(".xml"):
            new = re.sub(rb"<v\s*/>", b"", data)
            new = re.sub(rb"<v></v>", b"", new)
            if new != data:
                items[name] = new
                changed = True
    if not changed:
        return
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)


def _coerce(v):
    v = (v or "").strip()
    if v.startswith("="):
        return v  # formula — keep as-is
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    try:
        return float(v)
    except ValueError:
        return v


def _need_openpyxl():
    if not _HAVE_OPENPYXL:
        return "openpyxl isn't installed — run: pip install openpyxl"
    return ""


def _alias(params):
    a = _aliases()
    name = (params.get("name") or "").strip().lower()
    if not name or (params.get("action") or "list").lower() == "list":
        if not a:
            return "No sheets registered yet. Say: remember sheet <nickname> <full path to .xlsx>"
        return "Registered sheets:\n" + "\n".join(
            f"  {k} -> {v}" for k, v in sorted(a.items()))
    path = os.path.expanduser((params.get("path") or "").strip().strip("\"'"))
    if (params.get("action") or "").lower() == "remove":
        a.pop(name, None)
        _save_aliases(a)
        return f"Forgot {name!r}."
    if not path.lower().endswith(".xlsx") or not os.path.isfile(path):
        return f"Can't register {name!r}: not an existing .xlsx file: {path or '(no path given)'}"
    a[name] = path
    _save_aliases(a)
    return f"Registered {name!r} -> {path}."


def _read(params):
    err = _need_openpyxl()
    if err:
        return err
    path, err = _resolve(params.get("file"))
    if err:
        return err
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[params["sheet"]] if params.get("sheet") else wb.active
    except KeyError:
        return f"Sheet {params.get('sheet')!r} not found in {os.path.basename(path)}."
    rng = (params.get("range") or "A1").strip().upper()
    try:
        if ":" in rng:
            rows = list(ws[rng])
        else:
            rows = [[ws[rng]]]
    except ValueError:
        return f"Bad range {rng!r} — use like A1 or A1:D10."
    lines = []
    for row in rows[:_MAX_ROWS]:
        lines.append(" | ".join("" if c.value is None else str(c.value) for c in row))
    if len(rows) > _MAX_ROWS:
        lines.append(f"... ({len(rows) - _MAX_ROWS} more rows)")
    return f"{os.path.basename(path)} [{ws.title}] {rng}:\n" + ("\n".join(lines) or "(empty)")


def _write(params):
    err = _need_openpyxl()
    if err:
        return err
    path, err = _resolve(params.get("file"))
    if err:
        return err
    cell = (params.get("cell") or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
        return f"Bad cell {cell!r} — use like B5."
    if params.get("value") is None:
        return "Nothing to write — tell me the value too."
    wb = openpyxl.load_workbook(path)
    ws = wb[params["sheet"]] if params.get("sheet") else wb.active
    ws[cell] = _coerce(params["value"])
    wb.save(path)
    _sanitize_xlsx(path)
    return f"Wrote {params['value']!r} to {cell} in {os.path.basename(path)} [{ws.title}]."


def _add_row(params):
    err = _need_openpyxl()
    if err:
        return err
    path, err = _resolve(params.get("file"))
    if err:
        return err
    raw = params.get("values") or ""
    values = next(csv.reader(io.StringIO(raw)), [])
    if not values:
        return "No values given — e.g. add row to budget: 2026-10-07, Metro, 108"
    wb = openpyxl.load_workbook(path)
    ws = wb[params["sheet"]] if params.get("sheet") else wb.active
    ws.append([_coerce(v) for v in values])
    wb.save(path)
    _sanitize_xlsx(path)
    return f"Added row {len(values)} cells to {os.path.basename(path)} [{ws.title}]."


TOOLS = [
    {"name": "excel_alias",
     "description": "Register a nickname for a workbook, list nicknames, or forget one. "
                    "Usage: remember sheet <nickname> <full path to .xlsx>.",
     "triggers": ["remember sheet", "register sheet", "sheet alias", "forget sheet",
                  "list sheets", "my sheets"],
     "parameters": {"type": "OBJECT", "properties": {
         "action": {"type": "STRING", "description": "add, list, or remove (default list)"},
         "name": {"type": "STRING", "description": "nickname, e.g. budget"},
         "path": {"type": "STRING", "description": "full path to the .xlsx file"}}},
     "arg_patterns": {"name": r"sheet\s+([A-Za-z0-9_\-]+)",
                      "path": r"([A-Za-z]:\\[^\s]+\.xlsx)"},
     "handler": _alias},
    {"name": "excel_read",
     "description": "Read cells from a workbook. Give a nickname or path, "
                    "optional sheet, and a cell or range like A1:D10.",
     "triggers": ["read excel", "excel read", "show excel", "check sheet",
                  "open sheet", "what's in excel"],
     "parameters": {"type": "OBJECT", "properties": {
         "file": {"type": "STRING", "description": "workbook nickname or .xlsx path"},
         "sheet": {"type": "STRING", "description": "sheet name (default: active sheet)"},
         "range": {"type": "STRING", "description": "cell or range, e.g. B5 or A1:D10"}}},
     "arg_patterns": {"range": r"\b([a-z]{1,3}[0-9]{1,5}(?::[a-z]{1,3}[0-9]{1,5})?)\b",
                      "file": [r"read excel\s+(\S+)", r"excel read\s+(\S+)",
                               r"show excel\s+(\S+)", r"check sheet\s+(\S+)",
                               r"open sheet\s+(\S+)",
                               r"what's in excel\s+(\S+)"]},
     "handler": _read},
    {"name": "excel_write",
     "description": "Write a value into a cell of a workbook. "
                    "Usage: update <nickname> cell B5 to 108.",
     "triggers": ["write excel", "update excel", "excel write", "set excel",
                  "change excel", "update"],
     "parameters": {"type": "OBJECT", "properties": {
         "file": {"type": "STRING", "description": "workbook nickname or .xlsx path"},
         "sheet": {"type": "STRING", "description": "sheet name (default: active sheet)"},
         "cell": {"type": "STRING", "description": "cell like B5"},
         "value": {"type": "STRING", "description": "value to write (numbers auto-detected, = starts a formula)"}}},
     "arg_patterns": {"cell": r"\b([a-z]{1,3}[1-9][0-9]*)\b",
                      "file": [r"update\s+(?:excel\s+)?(\S+)\s+cell",
                               r"write\s+(?:excel\s+)?(\S+)\s+cell",
                               r"set\s+(?:excel\s+)?(\S+)\s+cell",
                               r"change\s+(?:excel\s+)?(\S+)\s+cell"],
                      "value": [r"cell\s+[a-z]{1,3}[1-9][0-9]*\s+to\s+(.+)$",
                                r"\bto\s+(.+)$"]},
     "handler": _write},
    {"name": "excel_add_row",
     "description": "Append a row of values to a workbook's sheet. "
                    "Usage: add row to <nickname>: 2026-10-07, Metro, 108",
     "triggers": ["add row", "append row", "new row", "add entry", "log expense"],
     "parameters": {"type": "OBJECT", "properties": {
         "file": {"type": "STRING", "description": "workbook nickname or .xlsx path"},
         "sheet": {"type": "STRING", "description": "sheet name (default: active sheet)"},
         "values": {"type": "STRING", "description": "comma-separated values for the row"}}},
     "arg_patterns": {"file": [r"add row to\s+(\S+?)\s*:",
                               r"append row to\s+(\S+?)\s*:",
                               r"new row in\s+(\S+?)\s*:",
                               r"add entry to\s+(\S+?)\s*:",
                               r"log expense (?:in|to)\s+(\S+?)\s*:"],
                      "values": r":\s*(.+)$"},
     "handler": _add_row},
]
